import asyncio
import json
import re
import threading
from io import BytesIO

import httpx
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.api import create_app
from app.config import Settings
from app.retrieval import SearchIndex
from tests.conftest import FakeProvider


def _files(questions: bytes, document: bytes, source_name: str = "source.json"):
    return {
        "questions_file": ("questions.json", questions, "application/json"),
        "document_file": (source_name, document, "application/octet-stream"),
    }


def test_endpoint_handles_multiple_questions_and_export(fake_provider: FakeProvider) -> None:
    client = TestClient(create_app(provider=fake_provider))
    questions = json.dumps(
        {"questions": ["Which cloud provider?", "Unknown backup region?"]}
    ).encode()
    response = client.post(
        "/api/v1/answer",
        files=_files(questions, b'{"hosting":{"provider":"GCP"}}'),
    )
    assert response.status_code == 200
    body = response.json()
    assert len(body["results"]) == 2
    assert body["results"][0]["status"] == "answered"
    assert body["results"][0]["citations"][0]["json_path"] == "$.hosting"
    assert body["results"][1]["status"] == "not_found"
    assert response.headers["x-request-id"] == body["request_id"]

    exported = client.post("/api/v1/export.xlsx", json={"results": body["results"]})
    assert exported.status_code == 200
    sheet = load_workbook(BytesIO(exported.content), read_only=True).active
    assert [cell.value for cell in sheet[1]] == [
        "id",
        "question",
        "answer",
        "comments",
        "confidence",
        "status",
        "citations",
        "error_code",
        "error_message",
    ]
    assert sheet[2][2].value == "Google Cloud Platform"


def test_endpoint_accepts_pdf_and_returns_page_citation(
    fake_provider: FakeProvider, pdf_bytes: bytes
) -> None:
    client = TestClient(create_app(provider=fake_provider))
    response = client.post(
        "/api/v1/answer",
        files=_files(b'["Which cloud provider?"]', pdf_bytes, "report.pdf"),
    )
    assert response.status_code == 200
    result = response.json()["results"][0]
    assert result["answer"] == "Google Cloud Platform"
    assert result["citations"][0]["page"] == 2


def test_endpoint_rejects_invalid_question_json_without_key(fake_provider: FakeProvider) -> None:
    client = TestClient(create_app(provider=fake_provider))
    response = client.post("/api/v1/answer", files=_files(b"{broken", b'{"cloud":"GCP"}'))
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "invalid_questions_json"


def test_endpoint_enforces_upload_limit(fake_provider: FakeProvider) -> None:
    client = TestClient(create_app(settings=Settings(max_document_bytes=8), provider=fake_provider))
    response = client.post("/api/v1/answer", files=_files(b'["Question?"]', b'{"cloud":"GCP"}'))
    assert response.status_code == 413
    assert response.json()["error"]["code"] == "document_too_large"


def test_export_escapes_spreadsheet_formulas(fake_provider: FakeProvider) -> None:
    client = TestClient(create_app(provider=fake_provider))
    payload = {
        "results": [
            {
                "id": "1",
                "question": '=HYPERLINK("bad")',
                "answer": "Data-Not-Found",
                "comments": "",
                "confidence": "low",
                "status": "not_found",
                "citations": [],
            }
        ]
    }
    response = client.post("/api/v1/export.xlsx", json=payload)
    sheet = load_workbook(BytesIO(response.content), read_only=True).active
    assert sheet[2][1].value.startswith("'=HYPERLINK")


@pytest.mark.asyncio
async def test_health_stays_responsive_during_slow_search(
    fake_provider: FakeProvider, monkeypatch
) -> None:
    started = threading.Event()
    release = threading.Event()
    original_search = SearchIndex.search

    def slow_search(self, question, embedding, top_k):
        started.set()
        release.wait(timeout=1)
        return original_search(self, question, embedding, top_k)

    monkeypatch.setattr(SearchIndex, "search", slow_search)
    app = create_app(provider=fake_provider)
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url="http://test"
    ) as client:
        answer_task = asyncio.create_task(
            client.post(
                "/api/v1/answer",
                files=_files(b'["Which cloud?"]', b'{"provider":"GCP"}'),
            )
        )
        try:
            assert await asyncio.to_thread(started.wait, 0.5)
            health = await asyncio.wait_for(client.get("/healthz"), timeout=0.3)
            assert health.status_code == 200
            assert not answer_task.done()
        finally:
            release.set()
            await answer_task


@pytest.mark.parametrize("source_name", ["source.json", "report.pdf"])
def test_endpoint_returns_and_exports_partial_results(source_name, pdf_bytes) -> None:
    from app.errors import ServiceError

    class MixedProvider(FakeProvider):
        async def answer(self, question, passages):
            if question == "Fail?":
                raise ServiceError(504, "model_timeout", "Answer service timed out.")
            return await super().answer(question, passages)

    document = pdf_bytes if source_name.endswith(".pdf") else b'{"provider":"GCP"}'
    with TestClient(create_app(provider=MixedProvider())) as client:
        response = client.post(
            "/api/v1/answer",
            files=_files(b'["Which cloud?", "Fail?", "Unknown backup?"]', document, source_name),
        )
        assert response.status_code == 200
        results = response.json()["results"]
        assert [result["status"] for result in results] == ["answered", "error", "not_found"]
        assert results[1]["answer"] is None and results[1]["confidence"] is None
        assert results[1]["error"]["code"] == "model_timeout"
        exported = client.post("/api/v1/export.xlsx", json={"results": results})
    assert exported.status_code == 200
    sheet = load_workbook(BytesIO(exported.content), read_only=True).active
    assert sheet[3][2].value is None
    assert sheet[3][5].value == "error"
    assert sheet[3][7].value == "model_timeout"


def test_export_rejects_error_result_with_invented_answer() -> None:
    payload = {
        "results": [
            {
                "id": "1",
                "question": "Cloud?",
                "answer": "Invented",
                "comments": "",
                "confidence": "high",
                "status": "error",
                "citations": [],
                "error": {"code": "model_timeout", "message": "Timed out"},
            }
        ]
    }
    with TestClient(create_app(provider=FakeProvider())) as client:
        assert client.post("/api/v1/export.xlsx", json=payload).status_code == 422


@pytest.mark.asyncio
@pytest.mark.parametrize("source_name", ["source.json", "report.pdf"])
async def test_full_endpoint_uses_langchain_with_mocked_openai(
    mock_openai_provider, source_name, pdf_bytes
) -> None:
    calls = []

    def respond(request):
        body = json.loads(request.content)
        calls.append((request.url.path, body))
        if request.url.path.endswith("/embeddings"):
            return httpx.Response(
                200,
                json={
                    "object": "list",
                    "model": "text-embedding-3-small",
                    "data": [
                        {
                            "object": "embedding",
                            "index": i,
                            "embedding": [float("GCP" in text), 1.0],
                        }
                        for i, text in enumerate(body["input"])
                    ],
                    "usage": {"prompt_tokens": 5, "total_tokens": 5},
                },
            )
        prompt = body["messages"][-1]["content"]
        passage_id = next(
            passage_id
            for passage_id, text in re.findall(r"\[([^\]]+)\] ([^\n]+)", prompt)
            if "GCP" in text
        )
        if "Fail?" in prompt:
            # An upstream success containing no choices must become a question
            # error, while the two other model responses survive in the same batch.
            choices = []
        else:
            supported = "Unknown backup?" not in prompt
            content = {
                "answer": "GCP" if supported else "Data-Not-Found",
                "comments": "Source checked.",
                "citation_ids": [passage_id] if supported else [],
                "is_answerable": supported,
                "evidence_strength": "high" if supported else "low",
                "evidence_quotes": [{"passage_id": passage_id, "quote": "GCP"}]
                if supported
                else [],
            }
            choices = [
                {
                    "index": 0,
                    "finish_reason": "stop",
                    "message": {"role": "assistant", "content": json.dumps(content)},
                }
            ]
        return httpx.Response(
            200,
            json={
                "id": "chatcmpl-test",
                "object": "chat.completion",
                "created": 0,
                "model": "gpt-4o-mini",
                "choices": choices,
            },
        )

    document = pdf_bytes if source_name.endswith(".pdf") else b'{"provider":"GCP"}'
    async with (
        mock_openai_provider(respond) as provider,
        httpx.AsyncClient(
            transport=httpx.ASGITransport(app=create_app(provider=provider)), base_url="http://test"
        ) as client,
    ):
        response = await client.post(
            "/api/v1/answer",
            files=_files(b'["Which cloud?", "Fail?", "Unknown backup?"]', document, source_name),
        )
    assert response.status_code == 200
    results = response.json()["results"]
    assert [result["status"] for result in results] == ["answered", "error", "not_found"]
    assert results[1]["error"]["code"] == "invalid_model_response"
    citation = results[0]["citations"][0]
    assert citation["source"] == source_name
    if source_name.endswith(".pdf"):
        assert citation["page"] == 2
    else:
        assert citation["json_path"] == "$"
    assert len([call for call in calls if call[0].endswith("/chat/completions")]) == 3


def test_invalid_source_json_never_reaches_embeddings(fake_provider) -> None:
    with TestClient(create_app(provider=fake_provider)) as client:
        response = client.post("/api/v1/answer", files=_files(b'["Cloud?"]', b'{"value":NaN}'))
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "invalid_source_json"
    assert fake_provider.embedding_batches == []
