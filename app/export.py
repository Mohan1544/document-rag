"""Create an Excel download from reviewed answer results."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.schemas import AnswerResult

HEADERS = ["id", "question", "answer", "comments", "confidence", "status", "citations"]


def _safe_cell(value: str) -> str:
    # Prevent spreadsheet formula execution when user supplied questions are exported.
    if value and value[0] in ("=", "+", "-", "@", "\t", "\r", "\n"):
        return "'" + value
    return value


def build_workbook(results: list[AnswerResult]) -> bytes:
    """Build a downloadable workbook from the results currently held by the UI."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Answers"
    sheet.append(HEADERS)
    for result in results:
        locations = []
        for citation in result.citations:
            location = f"{citation.source}"
            if citation.page is not None:
                location += f" page {citation.page}"
            if citation.json_path is not None:
                location += f" {citation.json_path}"
            locations.append(location)
        sheet.append(
            [
                _safe_cell(result.id),
                _safe_cell(result.question),
                _safe_cell(result.answer),
                _safe_cell(result.comments),
                result.confidence,
                result.status,
                _safe_cell("; ".join(locations)),
            ]
        )
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = [18, 55, 70, 70, 15, 17, 45]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="17324D")
        cell.font = Font(color="FFFFFF", bold=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        sheet.row_dimensions[row[0].row].height = 52
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
